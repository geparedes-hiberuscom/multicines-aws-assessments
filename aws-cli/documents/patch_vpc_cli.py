"""
Add VPC CLI box to section 4 of Manual de Gestión de Recursos en AWS v1.0.docx
Inserts before section 5 (Route 53)
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

DOCX_PATH = os.path.join(os.path.dirname(__file__), "Manual de Gestión de Recursos en AWS v1.0.docx")
doc = Document(DOCX_PATH)
body = doc.element.body

# Find section 5 heading to insert before it
h5_elem = None
for p in doc.paragraphs:
    if p.style.name == 'Heading 1' and 'Route 53' in p.text:
        h5_elem = p._element
        break

if h5_elem is None:
    print("[ERROR] Section 5 not found")
    exit(1)

# Add CLI label + table at end, then move before section 5
label = doc.add_paragraph()
run = label.add_run("âŒ¨ AWS CLI")
run.bold = True
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

# CLI table
cli_text = """# === CONSULTA ===
# Listar VPCs
aws ec2 describe-vpcs --region us-east-1 --query "Vpcs[].{ID:VpcId,CIDR:CidrBlock,Name:Tags[?Key=='Name'].Value|[0]}" --output table

# Listar subnets de una VPC
aws ec2 describe-subnets --filters "Name=vpc-id,Values=vpc-0a5ebea8e7bee9ed5" --region us-east-1 --query "Subnets[].{ID:SubnetId,CIDR:CidrBlock,AZ:AvailabilityZone,Name:Tags[?Key=='Name'].Value|[0]}" --output table

# Listar Security Groups
aws ec2 describe-security-groups --filters "Name=vpc-id,Values=vpc-0a5ebea8e7bee9ed5" --region us-east-1 --query "SecurityGroups[].{ID:GroupId,Name:GroupName}" --output table

# Ver reglas de un SG
aws ec2 describe-security-groups --group-ids sg-0fdf0c744482646ae --region us-east-1

# Listar route tables
aws ec2 describe-route-tables --filters "Name=vpc-id,Values=vpc-0a5ebea8e7bee9ed5" --region us-east-1 --query "RouteTables[].{ID:RouteTableId,Name:Tags[?Key=='Name'].Value|[0]}" --output table

# === CREACIÃ“N ===
# Crear VPC
aws ec2 create-vpc --cidr-block 10.0.0.0/16 --tag-specifications 'ResourceType=vpc,Tags=[{Key=Name,Value=vpc-multicines-dev}]' --region us-east-1

# Crear subnet
aws ec2 create-subnet --vpc-id <VPC_ID> --cidr-block 10.0.1.0/24 --availability-zone us-east-1a --tag-specifications 'ResourceType=subnet,Tags=[{Key=Name,Value=subnet-dev-public-1a}]' --region us-east-1

# Crear Internet Gateway + attach
aws ec2 create-internet-gateway --tag-specifications 'ResourceType=internet-gateway,Tags=[{Key=Name,Value=dev-igw}]' --region us-east-1
aws ec2 attach-internet-gateway --internet-gateway-id <IGW_ID> --vpc-id <VPC_ID> --region us-east-1

# Crear NAT Gateway
aws ec2 allocate-address --domain vpc --region us-east-1
aws ec2 create-nat-gateway --subnet-id <PUBLIC_SUBNET_ID> --allocation-id <EIP_ID> --tag-specifications 'ResourceType=natgateway,Tags=[{Key=Name,Value=nat-dev}]' --region us-east-1

# Crear Security Group
aws ec2 create-security-group --group-name sg-alb-dev --description "ALB Security Group" --vpc-id <VPC_ID> --region us-east-1
aws ec2 authorize-security-group-ingress --group-id <SG_ID> --protocol tcp --port 443 --cidr 0.0.0.0/0 --region us-east-1"""

t = doc.add_table(rows=1, cols=1)
t.style = 'Table Grid'
cell = t.rows[0].cells[0]
cell.text = ''
tc = cell._tc
tcPr = tc.get_or_add_tcPr()
tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2" w:val="clear"/>'))

for line in cli_text.strip().split('\n'):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(line)
    run.font.name = 'Consolas'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

# Remove first empty paragraph in cell
if cell.paragraphs[0].text == '':
    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)

spacer = doc.add_paragraph()

# Move label + table + spacer before section 5
label_elem = label._element
tbl_elem = t._tbl
spacer_elem = spacer._element

body.remove(label_elem)
body.remove(tbl_elem)
body.remove(spacer_elem)

h5_elem.addprevious(spacer_elem)
h5_elem.addprevious(tbl_elem)
h5_elem.addprevious(label_elem)

doc.save(DOCX_PATH)
print(f"[DONE] VPC CLI box added before section 5")
print(f"[DONE] Size: {os.path.getsize(DOCX_PATH) / 1024:.0f} KB")
