"""
Final fix: Remove ALL VPN-related content (duplicates from previous runs),
then insert a clean section 13 before Anexo.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement
import os

DOCX_PATH = os.path.join(os.path.dirname(__file__), "Manual Multicines v2.docx")
doc = Document(DOCX_PATH)
body = doc.element.body

# === Step 1: Find CI/CD section end and Anexo start ===
cicd_heading = None
anexo_heading = None

for p in doc.paragraphs:
    if p.style.name == 'Heading 1':
        if 'CI/CD' in p.text:
            cicd_heading = p._element
        if 'Anexo' in p.text:
            anexo_heading = p._element

if not anexo_heading:
    print("[ERROR] Anexo heading not found")
    exit(1)

print(f"[INFO] Found CI/CD heading: {cicd_heading is not None}")
print(f"[INFO] Found Anexo heading: {anexo_heading is not None}")

# === Step 2: Remove everything between end of CI/CD content and Anexo ===
# Find all elements between CI/CD's last sub-element and Anexo
# Strategy: collect elements that are after CI/CD Heading 2 content ends
# and before Anexo

# First, find the last Heading 2 that belongs to CI/CD (12.x)
all_elements = list(body)
cicd_idx = all_elements.index(cicd_heading)
anexo_idx = all_elements.index(anexo_heading)

# Everything between CI/CD section's legitimate content and Anexo is VPN junk
# CI/CD section has sub-headings like "12.1", "12.2", etc.
# Find the last element that legitimately belongs to CI/CD
# (the protection de ramas images end the section)

# Find last element before any VPN content starts
# VPN content starts after the branch protection images
# Let's find elements between CI/CD heading and Anexo that contain "VPN", "Customer Gateway", etc.

to_remove = []
in_vpn_zone = False
for i in range(cicd_idx + 1, anexo_idx):
    elem = all_elements[i]
    # Check if this element has VPN-related text
    text = ''.join(t.text or '' for t in elem.findall(f'.//{qn("w:t")}'))
    
    # Check for heading style
    pStyle_elem = elem.find(f'.//{qn("w:pStyle")}')
    style_val = pStyle_elem.get(qn('w:val'), '') if pStyle_elem is not None else ''
    
    # If it's a Heading1 or Heading2 with VPN-related text, mark zone
    if 'Heading' in style_val:
        vpn_keywords = ['VPN', 'Customer Gateway', 'Virtual Private Gateway', 
                       'Route Propagation', 'Parámetros de configuración',
                       'Verificación', 'Troubleshooting', '13.']
        if any(kw in text for kw in vpn_keywords):
            in_vpn_zone = True
    
    if in_vpn_zone:
        to_remove.append(elem)

# Also check for page breaks and empty paragraphs just before VPN content
# that are orphaned
if to_remove:
    # Look backwards from first VPN element for orphaned page breaks
    first_vpn_idx = all_elements.index(to_remove[0])
    for i in range(first_vpn_idx - 1, cicd_idx, -1):
        elem = all_elements[i]
        text = ''.join(t.text or '' for t in elem.findall(f'.//{qn("w:t")}'))
        has_page_break = elem.find(f'.//{qn("w:br")}') is not None
        if has_page_break or (not text.strip()):
            to_remove.insert(0, elem)
        else:
            break

print(f"[INFO] Removing {len(to_remove)} VPN/orphan elements")
for elem in to_remove:
    body.remove(elem)

# === Step 3: Rename Anexo to 14 ===
for p in doc.paragraphs:
    if p.style.name == 'Heading 1' and 'Anexo' in p.text:
        for run in p.runs:
            run.text = ''
        if p.runs:
            p.runs[0].text = '14. Anexo — Resumen de Recursos por Ambiente'
        print("[INFO] Renamed Anexo to 14")
        anexo_heading = p._element
        break

# === Step 4: Insert clean VPN section 13 ===
def mk_h(text, level):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), f'Heading{level}')
    pPr.append(pStyle)
    p.append(pPr)
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    p.append(r)
    return p

def mk_p(text, bold=False, indent=False, color=None):
    p = OxmlElement('w:p')
    if indent:
        pPr = OxmlElement('w:pPr')
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '567')
        pPr.append(ind)
        p.append(pPr)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '20')
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    p.append(r)
    return p

def mk_br():
    p = OxmlElement('w:p')
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    p.append(r)
    return p

elements = []
elements.append(mk_br())
elements.append(mk_h('13. VPN Site-to-Site', 1))
elements.append(mk_p('Se establece una conexión VPN Site-to-Site por ambiente para comunicar la VPC en AWS con la red on-premise de Multicines. Se configura un Customer Gateway (CGW) con la IP pública del router on-premise, un Virtual Private Gateway (VGW) adjunto a la VPC, y una VPN Connection tipo IPSec que los asocia.'))
elements.append(mk_p(''))

elements.append(mk_h('13.1 Customer Gateway', 2))
steps = [
    '1. VPC → panel izquierdo → Customer Gateways → Create Customer Gateway',
    '2. Name: {env}-multicines-integration-cgw',
    '3. Routing: Static',
    '4. IP Address: IP pública del router on-premise (proporcionada por Multicines)',
    '5. BGP ASN: 65000',
    '6. Click Create Customer Gateway',
]
for s in steps:
    elements.append(mk_p(s, indent=True))
elements.append(mk_p(''))
elements.append(mk_p('[TODO: Insertar pantalla del formulario de creación de Customer Gateway]', bold=True, color='FF0000'))
elements.append(mk_p(''))

elements.append(mk_h('13.2 Virtual Private Gateway', 2))
steps = [
    '1. VPC → Virtual Private Gateways → Create Virtual Private Gateway',
    '2. Name: {env}-multicines-integration-vgw',
    '3. ASN: Amazon default ASN',
    '4. Click Create Virtual Private Gateway',
    '5. Seleccionar el VGW → Actions → Attach to VPC',
    '6. Seleccionar la VPC del ambiente correspondiente',
    '7. Click Attach to VPC',
]
for s in steps:
    elements.append(mk_p(s, indent=True))
elements.append(mk_p(''))
elements.append(mk_p('[TODO: Insertar pantalla del VGW adjunto a la VPC]', bold=True, color='FF0000'))
elements.append(mk_p(''))

elements.append(mk_h('13.3 VPN Connection', 2))
steps = [
    '1. VPC → Site-to-Site VPN Connections → Create VPN Connection',
    '2. Name: {env}-multicines-integration-vpn',
    '3. Target Gateway Type: Virtual Private Gateway → seleccionar VGW creado',
    '4. Customer Gateway: Existing → seleccionar CGW creado',
    '5. Routing Options: Static',
    '6. Click Create VPN Connection',
    '7. Seleccionar la conexión → Download Configuration para config del router',
]
for s in steps:
    elements.append(mk_p(s, indent=True))
elements.append(mk_p(''))
elements.append(mk_p('[TODO: Insertar pantalla del wizard de creación de VPN Connection]', bold=True, color='FF0000'))
elements.append(mk_p(''))

elements.append(mk_h('13.4 Route Propagation', 2))
steps = [
    '1. VPC → Route Tables → seleccionar tabla de rutas privada del ambiente',
    '2. Pestaña Route propagation → Edit route propagation',
    '3. Habilitar propagation para el VGW creado (checkbox)',
    '4. Click Save',
    '5. Verificar en pestaña Routes que aparezcan rutas hacia la red on-premise',
]
for s in steps:
    elements.append(mk_p(s, indent=True))
elements.append(mk_p(''))
elements.append(mk_p('[TODO: Insertar pantalla de Route propagation habilitado]', bold=True, color='FF0000'))
elements.append(mk_p(''))

elements.append(mk_h('13.5 Verificación', 2))
steps = [
    '1. VPN Connections → Tunnel Details: al menos un túnel con estado UP',
    '2. Route Tables privadas: verificar rutas propagadas hacia CIDR on-premise',
    '3. CloudWatch → Métricas → AWS/VPN → TunnelState (1=UP, 0=DOWN)',
]
for s in steps:
    elements.append(mk_p(s, indent=True))
elements.append(mk_p(''))

# Insert before Anexo (in correct order - first to last)
for elem in elements:
    anexo_heading.addprevious(elem)

print(f"[INFO] Inserted {len(elements)} clean VPN elements before Anexo")

doc.save(DOCX_PATH)
print(f"\n[DONE] {DOCX_PATH}")
print(f"[DONE] Size: {os.path.getsize(DOCX_PATH) / 1024:.0f} KB")
