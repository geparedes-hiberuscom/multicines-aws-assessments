"""
Update section 10.4 Variables de Entorno table with real values.
Then apply image borders.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

DOCX_PATH = os.path.join(os.path.dirname(__file__), "Manual Multicines v3.docx")
doc = Document(DOCX_PATH)
body = doc.element.body

# Find the env vars table (between 10.4 and 10.5)
all_elems = list(body)
h104 = doc.paragraphs[303]._element
h105 = doc.paragraphs[307]._element
h104_idx = all_elems.index(h104)
h105_idx = all_elems.index(h105)

target_table = None
for t in doc.tables:
    t_idx = all_elems.index(t._tbl)
    if h104_idx < t_idx < h105_idx:
        target_table = t
        break

if target_table is None:
    print("[ERROR] Table not found")
    exit(1)

# Clear all rows except header
while len(target_table.rows) > 1:
    tr = target_table.rows[-1]._tr
    target_table._tbl.remove(tr)

# Update header style
headers = ["Variable", "Dev", "Prod"]
for i, h in enumerate(headers):
    cell = target_table.rows[0].cells[i]
    cell.text = ''
    r = cell.paragraphs[0].add_run(h)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    existing = tcPr.find(qn('w:shd'))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F4E79" w:val="clear"/>'))

# Add all environment variables
env_vars = [
    ["SPRING_PROFILES_ACTIVE", "dev", "prod"],
    ["AWS_REGION", "us-east-1", "us-east-1"],
    ["OTEL_LOGS_EXPORTER", "none", "none"],
    ["OTEL_METRICS_EXPORTER", "none", "none"],
    ["OTEL_TRACES_EXPORTER", "otlp", "otlp"],
    ["OTEL_SERVICE_NAME", "integrator", "integrator"],
    ["OTEL_EXPORTER_OTLP_ENDPOINT", "http://localhost:4317", "http://localhost:4317"],
    ["OTEL_EXPORTER_OTLP_PROTOCOL", "grpc", "grpc"],
    ["OTEL_TRACES_SAMPLER", "parentbased_traceidratio", "parentbased_traceidratio"],
    ["OTEL_TRACES_SAMPLER_ARG", "1.0 (100%)", "0.10 (10%)"],
]

for row_data in env_vars:
    row = target_table.add_row()
    for i, val in enumerate(row_data):
        cell = row.cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(str(val))
        r.font.size = Pt(9)

print(f"[OK] Updated env vars table: {len(env_vars)} rows")

# === Apply image borders ===
count = 0
for p in doc.paragraphs:
    for run in p.runs:
        drawings = run._element.findall('.//' + qn('wp:inline'))
        for drawing in drawings:
            pics = drawing.findall('.//' + qn('pic:pic'))
            for pic in pics:
                spPr = pic.find(qn('pic:spPr'))
                if spPr is None:
                    continue
                existing = spPr.find(qn('a:ln'))
                if existing is not None:
                    spPr.remove(existing)
                ln = parse_xml(
                    '<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="12700">'
                    '<a:solidFill><a:srgbClr val="808080"/></a:solidFill>'
                    '</a:ln>'
                )
                spPr.append(ln)
                count += 1

print(f"[OK] Image borders applied: {count} images")

doc.save(DOCX_PATH)
print(f"\n[DONE] {DOCX_PATH}")
print(f"[DONE] Size: {os.path.getsize(DOCX_PATH) / 1024:.0f} KB")
