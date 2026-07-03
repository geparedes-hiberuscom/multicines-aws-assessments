"""
Build Manual de Gesti�n de Recursos en AWS v1.0.docx - Professional version
- Copies original DOCX (keeps cover page + TOC)
- Removes everything from "1. Introducción" onwards
- Rebuilds all sections with consistent professional format
- Each section: description + console steps + params table + CLI box + verification
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from docx.enum.table import WD_TABLE_ALIGNMENT
import os, shutil

BASE_DIR = os.path.dirname(__file__)
SOURCE = os.path.join(BASE_DIR, "..", "..", "Insumos iniciales", "Manual para multicines.docx")
OUTPUT = os.path.join(BASE_DIR, "Manual de Gesti�n de Recursos en AWS v1.0.docx")
IMAGES = os.path.join(BASE_DIR, "images", "media")

shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)

# === Remove everything from "1. Introducción" onwards ===
body = doc.element.body
intro_found = False
to_remove = []
for elem in list(body):
    if elem.tag.endswith('}sectPr'):
        continue
    if not intro_found:
        texts = ''.join(t.text or '' for t in elem.findall(f'.//{qn("w:t")}'))
        if '1. Introducci' in texts or 'Introducción' in texts:
            pStyle = elem.find(f'.//{qn("w:pStyle")}')
            if pStyle is not None and 'Heading' in pStyle.get(qn('w:val'), ''):
                intro_found = True
    if intro_found:
        to_remove.append(elem)

for elem in to_remove:
    body.remove(elem)
print(f"[INFO] Removed {len(to_remove)} elements from Introducción onwards")

# === Helper functions ===
def h1(text):
    doc.add_heading(text, level=1)

def h2(text):
    doc.add_heading(text, level=2)

def h3(text):
    doc.add_heading(text, level=3)

def para(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    return p

def bold_para(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    return p

def todo(text):
    p = doc.add_paragraph()
    run = p.add_run(f"[TODO: {text}]")
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    run.font.size = Pt(9)

def steps(items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(f"{i}. {item}")
        run.font.size = Pt(10)

def table(headers, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Blue header background
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F4E79" w:val="clear"/>'))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri+1].cells[ci]
            cell.text = ''
            r = cell.paragraphs[0].add_run(str(val))
            r.font.size = Pt(9)
    doc.add_paragraph()

def cli_box(text):
    """Professional CLI code block: 1-cell table with grey bg"""
    bold_para("⌨ AWS CLI")
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    cell = t.rows[0].cells[0]
    cell.text = ''
    # Grey background
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2" w:val="clear"/>'))
    # Add text lines
    for line in text.strip().split('\n'):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    # Remove first empty paragraph
    if cell.paragraphs[0].text == '':
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    doc.add_paragraph()

def add_image(filename):
    """Add image if exists, otherwise TODO"""
    path = os.path.join(IMAGES, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(5.5))
    else:
        todo(f"Insertar pantalla: {filename}")

def page_break():
    doc.add_page_break()

# =====================================================================
# SECTION 1: INTRODUCCIÓN
# =====================================================================
h1("1. Introducción")
h2("1.1 Propósito")
para("Este documento describe la infraestructura cloud implementada en AWS para el proyecto Multicines. Detalla cada componente, las decisiones técnicas y los pasos para la administración operativa.")
h2("1.2 Contexto")
para("Multicines expone su sistema Vista API Connect hacia servicios externos a través de una capa de integración en AWS. La solución contempla dos ambientes (dev/prod) con despliegues automatizados via GitHub Actions. El acceso público se realiza a través del dominio cloudmulticines.com.")
h2("1.3 Datos de la cuenta")
table(["Parámetro", "Valor"], [
    ["Cuenta AWS", "340271092920"],
    ["Región", "us-east-1 (N. Virginia)"],
    ["Dominio", "cloudmulticines.com"],
    ["Certificado", "*.cloudmulticines.com (ISSUED)"],
    ["ECS Cluster", "multicines-cluster"],
    ["ECR", "multicines/integration-service"],
])

h2("1.4 Inventario de Recursos Existentes")
h3("Recursos Compartidos (shared)")
table(["Servicio", "Recurso", "ID/Nombre"], [
    ["Route53", "Hosted Zone", "cloudmulticines.com (Z02656981PPWUEVYOEMVL)"],
    ["ACM", "Certificado SSL", "*.cloudmulticines.com (ISSUED)"],
    ["ECR", "Repositorio", "multicines/integration-service"],
    ["ECS", "Cluster", "multicines-cluster"],
    ["IAM", "Execution Role", "ecsTaskExecutionRole"],
    ["IAM", "Task Role", "multicines-ecs-task-role"],
    ["IAM", "GitHub Actions Role", "multicines-github-actions-role"],
    ["EC2", "Customer Gateway", "cgw-090efcb086447cd1a (IP: 200.7.217.58)"],
])
h3("Ambiente Dev")
table(["Servicio", "Recurso", "ID/Nombre"], [
    ["VPC", "VPC", "vpc-0a5ebea8e7bee9ed5"],
    ["VPC", "Subnet pública A", "subnet-0c92479c2831f04eb"],
    ["VPC", "Subnet pública B", "subnet-02193c97ceb6ec25d"],
    ["VPC", "Subnet privada", "subnet-0eeec60453f7f9492"],
    ["EC2", "SG ALB", "sg-0fdf0c744482646ae"],
    ["EC2", "SG ECS", "sg-01c619444bab0b8d3"],
    ["ALB", "Load Balancer", "dev-multicines-integration-alb"],
    ["ALB", "Target Group", "dev-multicines-integration-tg (port 8095)"],
    ["Route53", "Record A", "api-dev.cloudmulticines.com → ALB dev"],
    ["ECS", "Service", "dev-multicines-integration-svc (1 task)"],
    ["ECS", "Task Definition", "dev-multicines-integration-task"],
    ["Secrets", "Secret", "platform/dev/resilience"],
    ["Secrets", "Secret", "integrator/dev/app"],
    ["SNS", "Topic", "dev-multicines-integration-alarms"],
    ["CloudWatch", "Dashboard", "dev-multicines-integration-dashboard"],
    ["CloudWatch", "Alarms", "7 alarmas (cpu, memory, 5xx, 4xx, unhealthy, response-time, tasks)"],
    ["VPN", "Connection", "vpn-08add07f10aaf549d (dev-multicines-integration-vpn)"],
    ["VPN", "Virtual Private GW", "vgw-0d4a9bbfacc384bdc (dev-multicines-integration-vgw)"],
])
h3("Ambiente Prod")
table(["Servicio", "Recurso", "ID/Nombre"], [
    ["VPC", "VPC", "vpc-0860a0d2cf4df6140"],
    ["VPC", "Subnet pública A", "subnet-04e4cd18763fd6f84"],
    ["VPC", "Subnet pública B", "subnet-05c92a4889326c68e"],
    ["VPC", "Subnet privada A", "subnet-0d55ebfe5adbd4e81"],
    ["VPC", "Subnet privada B", "subnet-04f8ac3335c426f57"],
    ["EC2", "SG ALB", "sg-04f21c4c1d065fe60"],
    ["EC2", "SG ECS", "sg-0b78c14cd02b2f6e1"],
    ["ALB", "Load Balancer", "prod-multicines-integration-alb"],
    ["ALB", "Target Group", "prod-multicines-integration-tg (port 8095)"],
    ["Route53", "Record A", "api.cloudmulticines.com → ALB prod"],
    ["ECS", "Service", "prod-multicines-integration-svc (2 tasks)"],
    ["ECS", "Task Definition", "prod-multicines-integration-task"],
    ["Secrets", "Secret", "platform/prod/resilience"],
    ["Secrets", "Secret", "integrator/prod/app"],
    ["SNS", "Topic", "prod-multicines-integration-alarms"],
    ["CloudWatch", "Dashboard", "prod-multicines-integration-dashboard"],
    ["VPN", "Connection", "vpn-03ad7d548894fe350 (prod-multicines-integration-vpn)"],
    ["VPN", "Virtual Private GW", "vgw-0954e2b1a370ae78e (prod-multicines-integration-vgw)"],
])

# =====================================================================
# SECTION 2: INGRESO A LA CONSOLA
# =====================================================================
page_break()
h1("2. Ingreso a la Consola de AWS")
para("URL: https://signin.aws.amazon.com | Account ID: 340271092920")
add_image("image.png")
h2("2.1 Ingreso con usuario Root")
steps(["Seleccionar Root user → ingresar email → Next", "Ingresar contraseña", "Ingresar código MFA", "Se muestra el dashboard de la consola"])
add_image("image2.png")
add_image("image3.png")
add_image("image4.png")
add_image("image5.png")
h2("2.2 Ingreso con usuario IAM")
steps(["Seleccionar IAM user → Account ID: 340271092920 → Next", "Ingresar usuario y contraseña → Sign in"])
add_image("image6.png")
add_image("image7.png")
add_image("image8.png")

# =====================================================================
# SECTION 3: IAM
# =====================================================================
page_break()
h1("3. Administración de Acceso (IAM)")
h2("3.1 Panel de IAM")
add_image("image9.png")
add_image("imagea.png")
h2("3.2 Grupos IAM")
add_image("imageb.png")
add_image("imagec.png")
add_image("imaged.png")
add_image("imagee.png")
add_image("imagef.png")
add_image("image10.png")
add_image("image11.png")
add_image("image12.png")
h2("3.3 Usuarios IAM")
add_image("image13.png")
add_image("image14.png")
add_image("image15.png")
add_image("image16.png")
add_image("image17.png")
add_image("image18.png")
add_image("image19.png")
add_image("image1a.png")
add_image("image1b.png")
add_image("image1c.png")
add_image("image1d.png")
h2("3.4 Políticas IAM")
add_image("image1e.png")
add_image("image1f.png")
cli_box("""# === GRUPOS ===
# Crear grupo
aws iam create-group --group-name Administrators

# Asignar política al grupo
aws iam attach-group-policy --group-name Administrators \\
  --policy-arn arn:aws:iam::aws:policy/AdministratorAccess

# Crear grupo read-only
aws iam create-group --group-name ReadOnlyUsers
aws iam attach-group-policy --group-name ReadOnlyUsers \\
  --policy-arn arn:aws:iam::aws:policy/ReadOnlyAccess

# Listar grupos
aws iam list-groups

# === USUARIOS ===
# Crear usuario con acceso a consola
aws iam create-user --user-name gustavo.paredes@hiberus.com
aws iam create-login-profile --user-name gustavo.paredes@hiberus.com \\
  --password <TEMP_PASSWORD> --password-reset-required

# Agregar usuario a grupo
aws iam add-user-to-group --user-name gustavo.paredes@hiberus.com --group-name Administrators

# Listar usuarios
aws iam list-users

# === POLÍTICAS ===
# Listar políticas adjuntas a un grupo
aws iam list-attached-group-policies --group-name Administrators

# Listar políticas adjuntas a un usuario
aws iam list-attached-user-policies --user-name gustavo.paredes@hiberus.com

# Asignar política directamente a usuario (no recomendado, usar grupos)
aws iam attach-user-policy --user-name <USER> --policy-arn <POLICY_ARN>""")
h2("3.5 Roles del Proyecto")
table(["Rol", "ARN", "Propósito"], [
    ["ecsTaskExecutionRole", "arn:aws:iam::340271092920:role/ecsTaskExecutionRole", "Ejecución ECS"],
    ["multicines-ecs-task-role", "arn:aws:iam::340271092920:role/multicines-ecs-task-role", "Task role + X-Ray"],
    ["multicines-github-actions-role", "arn:aws:iam::340271092920:role/multicines-github-actions-role", "CI/CD OIDC"],
])
cli_box("""# Crear rol de ejecución ECS
aws iam create-role --role-name ecsTaskExecutionRole \\
  --assume-role-policy-document file://ecs-trust-policy.json

aws iam attach-role-policy --role-name ecsTaskExecutionRole \\
  --policy-arn arn:aws:iam::aws:policy/service-role/AmazonECSTaskExecutionRolePolicy

# Crear task role con X-Ray
aws iam create-role --role-name multicines-ecs-task-role \\
  --assume-role-policy-document file://ecs-trust-policy.json

aws iam put-role-policy --role-name multicines-ecs-task-role \\
  --policy-name dev-multicines-integration-xray-policy \\
  --policy-document file://xray-policy.json""")

# =====================================================================
# SECTION 4: VPC
# =====================================================================
page_break()
h1("4. Administración de VPC")
h2("4.1 VPCs del proyecto")
table(["VPC", "ID", "Ambiente"], [
    ["vpc-multicines-dev", "vpc-0a5ebea8e7bee9ed5", "dev"],
    ["prod-vpc-multicines", "vpc-0860a0d2cf4df6140", "prod"],
])
add_image("image20.png")
add_image("image21.png")
add_image("image22.png")
add_image("image23.png")
add_image("image24.png")
add_image("image25.png")
add_image("image26.png")
add_image("image27.png")
add_image("image28.png")
add_image("image29.png")
add_image("image2a.png")
h2("4.2 Internet Gateway")
add_image("image2b.png")
add_image("image2c.png")
h2("4.3 NAT Gateway")
add_image("image2d.png")
add_image("image2e.png")
add_image("image2f.png")
h2("4.4 Security Groups")
table(["Security Group", "ID", "Ambiente"], [
    ["SG ALB", "sg-0fdf0c744482646ae", "dev"],
    ["SG ECS", "sg-01c619444bab0b8d3", "dev"],
    ["SG ALB", "sg-04f21c4c1d065fe60", "prod"],
    ["SG ECS", "sg-0b78c14cd02b2f6e1", "prod"],
])
add_image("image30.png")
add_image("image31.png")
add_image("image32.png")
add_image("image33.png")
add_image("image34.png")
h2("4.5 Tablas de Enrutamiento")
add_image("image35.png")
add_image("image36.png")
add_image("image37.png")
add_image("image38.png")
add_image("image39.png")

# =====================================================================
# SECTION 5: ROUTE 53
# =====================================================================
page_break()
h1("5. Route 53 — DNS")
para("Dominio: cloudmulticines.com | Zone ID: Z02656981PPWUEVYOEMVL")
h2("5.1 Registro del dominio")
bold_para("Opción A: Registrar dominio en Route53 (recomendado)")
steps(["Route 53 → Registered domains → Register domains", "Buscar: cloudmulticines.com → Select → Proceed to checkout", "Llenar datos de contacto → aceptar términos → Submit", "Costo: ~$12/año. AWS auto-crea la hosted zone y configura NS."])
todo("Insertar pantalla de Registered domains en Route53")
bold_para("Opción B: Crear hosted zone manualmente (dominio externo)")
steps(["Route 53 → Hosted zones → Create hosted zone", "Domain name: cloudmulticines.com → Type: Public → Create", "Copiar los 4 nameservers y configurarlos en el registrador externo"])
cli_box("""# Opción B: Crear hosted zone manual
aws route53 create-hosted-zone --name cloudmulticines.com \\
  --caller-reference "create-$(date +%s)"

# Ver nameservers de la zona
aws route53 get-hosted-zone --id Z02656981PPWUEVYOEMVL \\
  --query "DelegationSet.NameServers"

# Listar hosted zones
aws route53 list-hosted-zones""")
h2("5.2 Registros DNS")
table(["Registro", "Tipo", "Target", "Ambiente"], [
    ["api-dev.cloudmulticines.com", "A (Alias)", "ALB dev", "dev"],
    ["api.cloudmulticines.com", "A (Alias)", "ALB prod", "prod"],
])
h3("Crear registro A alias (consola)")
steps(["Route 53 → Hosted zones → cloudmulticines.com → Create record", "Record name: api-dev (o api para prod)", "Record type: A → Toggle Alias: ON", "Route traffic to: Application Load Balancer → us-east-1 → seleccionar ALB", "Evaluate target health: Yes → Create records"])
todo("Insertar pantalla del formulario de creación de record alias A")
cli_box("""# Crear registro A alias (dev)
aws route53 change-resource-record-sets \\
  --hosted-zone-id Z02656981PPWUEVYOEMVL \\
  --change-batch file://route53-dev-record.json

# Ejemplo route53-dev-record.json:
# {"Changes":[{"Action":"UPSERT","ResourceRecordSet":{
#   "Name":"api-dev.cloudmulticines.com","Type":"A",
#   "AliasTarget":{"HostedZoneId":"<ALB_ZONE>",
#   "DNSName":"dev-multicines-integration-alb-1885310298.us-east-1.elb.amazonaws.com",
#   "EvaluateTargetHealth":true}}}]}

# Listar records
aws route53 list-resource-record-sets --hosted-zone-id Z02656981PPWUEVYOEMVL

# Eliminar record
# Usar Action: "DELETE" en el change-batch con mismos valores""")

# =====================================================================
# SECTION 6: ACM
# =====================================================================
page_break()
h1("6. Certificados SSL — ACM")
para("Certificado wildcard *.cloudmulticines.com solicitado via DNS validation en ACM. Se auto-renueva mientras el CNAME de validación exista en la hosted zone.")
h2("6.1 Solicitar certificado (DNS validation)")
table(["Parámetro", "Valor"], [
    ["Dominio", "*.cloudmulticines.com"],
    ["Método", "DNS validation (auto-renovable)"],
    ["Estado", "ISSUED"],
    ["ARN", "arn:aws:acm:us-east-1:340271092920:certificate/eb5ce963-b419-4672-bb8a-7ed97becfafc"],
    ["Tags", "Name=multicines-integration-acm, Environment=shared"],
])
h3("Paso a paso (consola)")
steps([
    "Certificate Manager → Request certificate → Request a public certificate → Next",
    "Domain names: *.cloudmulticines.com",
    "Validation method: DNS validation",
    "Key algorithm: RSA 2048",
    "Tags: Name=multicines-integration-acm, Project=multicines, Environment=shared",
    "Click Request",
    "En el certificado → Create records in Route 53 (crea CNAME automáticamente)",
    "Esperar 2-5 minutos → estado cambia a ISSUED",
])
todo("Insertar pantalla del certificado en ACM con estado ISSUED")
todo("Insertar pantalla del CNAME de validación en la hosted zone")
h2("6.2 Renovación")
para("El certificado se auto-renueva ~60 días antes de expirar. Requisitos: que el CNAME de validación exista y que el cert esté asociado a un recurso (ALB). No requiere acción manual.")
cli_box("""# Solicitar certificado wildcard
aws acm request-certificate \\
  --domain-name "*.cloudmulticines.com" \\
  --validation-method DNS \\
  --tags Key=Name,Value=multicines-integration-acm Key=Project,Value=multicines Key=Environment,Value=shared \\
  --region us-east-1

# Ver estado del certificado
aws acm describe-certificate \\
  --certificate-arn arn:aws:acm:us-east-1:340271092920:certificate/eb5ce963-b419-4672-bb8a-7ed97becfafc \\
  --region us-east-1 --query "Certificate.{Status:Status,Expiry:NotAfter}"

# Listar certificados
aws acm list-certificates --region us-east-1

# Eliminar certificado (solo si no está asociado a ALB)
aws acm delete-certificate --certificate-arn <ARN> --region us-east-1""")

# =====================================================================
# SECTION 7: ALB
# =====================================================================
page_break()
h1("7. Application Load Balancer")
para("El ALB distribuye tráfico HTTPS al Target Group con las tareas ECS. Incluye listener HTTPS:443 con TLS 1.3 y redirect HTTP:80→HTTPS.")
h2("7.1 ALBs existentes")
table(["ALB", "DNS", "Ambiente"], [
    ["dev-multicines-integration-alb", "dev-multicines-integration-alb-1885310298.us-east-1.elb.amazonaws.com", "dev"],
    ["prod-multicines-integration-alb", "prod-multicines-integration-alb-43560610.us-east-1.elb.amazonaws.com", "prod"],
])
h2("7.2 Configuración")
table(["Parámetro", "Dev", "Prod"], [
    ["Tipo", "application, internet-facing", "application, internet-facing"],
    ["SSL Policy", "ELBSecurityPolicy-TLS13-1-2-2021-06", "ELBSecurityPolicy-TLS13-1-2-2021-06"],
    ["Certificado", "*.cloudmulticines.com", "*.cloudmulticines.com"],
    ["Listener HTTPS", "443 → Forward to TG", "443 → Forward to TG"],
    ["Listener HTTP", "80 → Redirect HTTPS 301", "80 → Redirect HTTPS 301"],
    ["Target Group", "dev-multicines-integration-tg", "prod-multicines-integration-tg"],
    ["TG Port", "8095", "8095"],
    ["Health Check", "/actuator/health:8095", "/actuator/health:8095"],
    ["Healthy/Unhealthy", "2 / 5", "2 / 5"],
])
h3("Crear ALB (consola)")
steps([
    "EC2 → Load Balancers → Create Load Balancer → Application Load Balancer",
    "Name: {env}-multicines-integration-alb, Scheme: Internet-facing",
    "Network mapping: VPC + 2 subnets públicas del ambiente",
    "Security groups: SG ALB del ambiente",
    "Listener HTTPS:443 → cert ACM *.cloudmulticines.com → Policy TLS 1.3",
    "Target Group: Create → IP type, port 8095, health check /actuator/health",
    "Agregar listener HTTP:80 → Redirect HTTPS 301",
])
todo("Insertar pantalla del ALB con listeners HTTPS y HTTP")
todo("Insertar pantalla del Target Group con targets healthy")
cli_box("""# Crear ALB (ejemplo dev)
aws elbv2 create-load-balancer --name dev-multicines-integration-alb \\
  --type application --scheme internet-facing \\
  --subnets subnet-0c92479c2831f04eb subnet-02193c97ceb6ec25d \\
  --security-groups sg-0fdf0c744482646ae \\
  --tags Key=Name,Value=dev-multicines-integration-alb Key=Project,Value=multicines Key=Environment,Value=dev \\
  --region us-east-1

# Crear Target Group
aws elbv2 create-target-group --name dev-multicines-integration-tg \\
  --protocol HTTP --port 8095 --vpc-id vpc-0a5ebea8e7bee9ed5 --target-type ip \\
  --health-check-path /actuator/health --health-check-port 8095 \\
  --health-check-interval-seconds 30 --health-check-timeout-seconds 10 \\
  --healthy-threshold-count 2 --unhealthy-threshold-count 5 --region us-east-1

# Crear Listener HTTPS:443
aws elbv2 create-listener --load-balancer-arn <ALB_ARN> \\
  --protocol HTTPS --port 443 --ssl-policy ELBSecurityPolicy-TLS13-1-2-2021-06 \\
  --certificates CertificateArn=arn:aws:acm:us-east-1:340271092920:certificate/eb5ce963-b419-4672-bb8a-7ed97becfafc \\
  --default-actions Type=forward,TargetGroupArn=<TG_ARN> --region us-east-1

# Crear Listener HTTP:80 redirect
aws elbv2 create-listener --load-balancer-arn <ALB_ARN> \\
  --protocol HTTP --port 80 \\
  --default-actions 'Type=redirect,RedirectConfig={Protocol=HTTPS,Port=443,StatusCode=HTTP_301}' --region us-east-1

# Eliminar: listeners → ALB → Target Group
aws elbv2 delete-listener --listener-arn <ARN> --region us-east-1
aws elbv2 delete-load-balancer --load-balancer-arn <ARN> --region us-east-1
aws elbv2 delete-target-group --target-group-arn <ARN> --region us-east-1""")

# =====================================================================
# SECTION 8: SECRETS MANAGER
# =====================================================================
page_break()
h1("8. Secrets Manager")
para("Almacena credenciales sensibles cifradas por ambiente. Los secretos son consumidos por las tareas ECS en tiempo de ejecución.")
h2("8.1 Secretos existentes")
table(["Nombre", "Ambiente"], [
    ["platform/dev/resilience", "dev"],
    ["integrator/dev/app", "dev"],
    ["platform/prod/resilience", "prod"],
    ["integrator/prod/app", "prod"],
])
h3("Crear secreto (consola)")
steps([
    "Secrets Manager → Store a new secret",
    "Secret type: Other type of secret",
    "Key/value: Plaintext → pegar JSON compactado",
    "Encryption key: aws/secretsmanager (default) → Next",
    "Secret name: ej. platform/dev/resilience",
    "Tags: Name, Project=multicines, Environment={env} → Next → Store",
])
todo("Insertar pantalla del listado de secretos")
cli_box("""# Crear secreto
aws secretsmanager create-secret --name platform/dev/resilience \\
  --secret-string '{"db_host":"...","db_pass":"..."}' \\
  --tags Key=Name,Value=platform/dev/resilience Key=Project,Value=multicines Key=Environment,Value=dev \\
  --region us-east-1

# Obtener valor
aws secretsmanager get-secret-value --secret-id platform/dev/resilience --region us-east-1

# Actualizar valor
aws secretsmanager update-secret --secret-id platform/dev/resilience \\
  --secret-string '{"db_host":"...","db_pass":"new"}' --region us-east-1

# Eliminar (sin período de recuperación)
aws secretsmanager delete-secret --secret-id platform/dev/resilience \\
  --force-delete-without-recovery --region us-east-1

# Listar secretos
aws secretsmanager list-secrets --region us-east-1""")

# =====================================================================
# SECTION 9: ECR
# =====================================================================
page_break()
h1("9. Elastic Container Registry (ECR)")
para("Repositorio privado: multicines/integration-service")
para("URI: 340271092920.dkr.ecr.us-east-1.amazonaws.com/multicines/integration-service")
add_image("image3a.png")
add_image("image3b.png")
add_image("image3c.png")
add_image("image3d.png")
add_image("image3e.png")
cli_box("""# Login a ECR
aws ecr get-login-password --region us-east-1 | \\
  docker login --username AWS --password-stdin 340271092920.dkr.ecr.us-east-1.amazonaws.com

# Crear repositorio
aws ecr create-repository --repository-name multicines/integration-service \\
  --image-tag-mutability IMMUTABLE --encryption-configuration encryptionType=AES256 \\
  --region us-east-1

# Build, tag y push
docker build -t multicines/integration-service:dev-latest .
docker tag multicines/integration-service:dev-latest \\
  340271092920.dkr.ecr.us-east-1.amazonaws.com/multicines/integration-service:dev-latest
docker push 340271092920.dkr.ecr.us-east-1.amazonaws.com/multicines/integration-service:dev-latest

# Listar imágenes
aws ecr list-images --repository-name multicines/integration-service --region us-east-1""")

# =====================================================================
# SECTION 10: ECS
# =====================================================================
page_break()
h1("10. Elastic Container Service (ECS)")
para("Cluster: multicines-cluster (Fargate). Servicios con sidecar ADOT para trazas X-Ray.")
h2("10.1 Cluster")
add_image("image3f.png")
add_image("image40.png")
add_image("image41.png")
add_image("image42.png")
add_image("image43.png")
h2("10.2 Servicios")
table(["Servicio", "Tasks", "Estado"], [
    ["dev-multicines-integration-svc", "1/1", "ACTIVE"],
    ["prod-multicines-integration-svc", "2/2", "ACTIVE"],
])
h2("10.3 Task Definition")
table(["Parámetro", "Dev", "Prod"], [
    ["Family", "dev-multicines-integration-task", "prod-multicines-integration-task"],
    ["CPU / Memoria", "512 / 1024 MB", "1024 / 2048 MB"],
    ["Container principal", "integration-service:8095", "integration-service:8095"],
    ["Sidecar ADOT", "aws-otel-collector:4317", "aws-otel-collector:4317"],
    ["Execution Role", "ecsTaskExecutionRole", "ecsTaskExecutionRole"],
    ["Task Role", "multicines-ecs-task-role", "multicines-ecs-task-role"],
    ["Log group", "/ecs/dev-multicines-integration", "/ecs/prod-multicines-integration"],
])
add_image("image44.png")
add_image("image45.png")
add_image("image46.png")
add_image("image47.png")
h2("10.4 Variables de Entorno")
table(["Variable", "Dev", "Prod"], [
    ["SPRING_PROFILES_ACTIVE", "dev", "prod"],
    ["OTEL_TRACES_EXPORTER", "otlp", "otlp"],
    ["OTEL_SERVICE_NAME", "integrator", "integrator"],
    ["OTEL_EXPORTER_OTLP_ENDPOINT", "http://localhost:4317", "http://localhost:4317"],
    ["OTEL_TRACES_SAMPLER", "parentbased_traceidratio", "parentbased_traceidratio"],
    ["OTEL_TRACES_SAMPLER_ARG", "1.0 (100%)", "0.10 (10%)"],
])
add_image("image48.png")
add_image("image49.png")
h2("10.5 Sidecar ADOT (X-Ray)")
table(["Parámetro", "Valor"], [
    ["Container", "aws-otel-collector"],
    ["Image", "public.ecr.aws/aws-observability/aws-otel-collector:latest"],
    ["Essential", "No"],
    ["Memory", "256 MB"],
    ["Puerto", "4317/tcp (gRPC OTLP)"],
])
add_image("image4a.png")
add_image("image4b.png")
todo("Insertar pantalla de Task Definition mostrando 2 containers")
cli_box("""# Crear cluster
aws ecs create-cluster --cluster-name multicines-cluster --region us-east-1

# Registrar Task Definition (desde JSON)
aws ecs register-task-definition --cli-input-json file://task-def.json --region us-east-1

# Crear servicio (dev)
aws ecs create-service --cluster multicines-cluster \\
  --service-name dev-multicines-integration-svc \\
  --task-definition dev-multicines-integration-task \\
  --launch-type FARGATE --desired-count 1 --enable-execute-command \\
  --network-configuration 'awsvpcConfiguration={subnets=["subnet-0eeec60453f7f9492"],securityGroups=["sg-01c619444bab0b8d3"],assignPublicIp=DISABLED}' \\
  --load-balancers targetGroupArn=<TG_ARN>,containerName=integration-service,containerPort=8095 \\
  --health-check-grace-period-seconds 210 --region us-east-1

# Force new deployment (reiniciar tareas con nuevos secrets/config)
aws ecs update-service --cluster multicines-cluster \\
  --service dev-multicines-integration-svc --force-new-deployment --region us-east-1

# Esperar estabilidad
aws ecs wait services-stable --cluster multicines-cluster \\
  --services dev-multicines-integration-svc --region us-east-1

# Eliminar servicio
aws ecs update-service --cluster multicines-cluster --service dev-multicines-integration-svc --desired-count 0 --region us-east-1
aws ecs delete-service --cluster multicines-cluster --service dev-multicines-integration-svc --region us-east-1""")

# =====================================================================
# SECTION 11: OBSERVABILIDAD
# =====================================================================
page_break()
h1("11. Observabilidad — CloudWatch, X-Ray")
para("Monitoreo completo: SNS para alertas, 7 alarmas CloudWatch, dashboard unificado, metric filters, y trazas X-Ray via sidecar ADOT.")
h2("11.1 SNS Topics")
table(["Topic", "Email", "Ambiente"], [
    ["dev-multicines-integration-alarms", "alarms@multicines.com.ec", "dev"],
    ["prod-multicines-integration-alarms", "alarms@multicines.com.ec", "prod"],
])
todo("Insertar pantalla del SNS Topic con suscripción confirmada")
h2("11.2 CloudWatch Alarms")
table(["Alarma", "Métrica", "Threshold", "Estado"], [
    ["alarm-cpu-high", "AWS/ECS CPUUtilization", "≥ 80%", "OK"],
    ["alarm-memory-high", "AWS/ECS MemoryUtilization", "≥ 80%", "OK"],
    ["alarm-5xx", "ALB HTTPCode_Target_5XX_Count", "≥ 10", "INSUFFICIENT_DATA"],
    ["alarm-4xx", "ALB HTTPCode_Target_4XX_Count", "≥ 50", "OK"],
    ["alarm-unhealthy-hosts", "ALB UnHealthyHostCount", "≥ 1", "OK"],
    ["alarm-response-time-high", "ALB TargetResponseTime", "≥ 5s/3s", "OK"],
    ["alarm-no-running-tasks", "ECS RunningTaskCount", "< 1", "OK"],
])
todo("Insertar pantalla del listado de alarmas en CloudWatch")
h2("11.3 CloudWatch Dashboard")
table(["Dashboard", "Ambiente"], [
    ["dev-multicines-integration-dashboard", "dev"],
    ["prod-multicines-integration-dashboard", "prod"],
])
para("Widgets: CPU, Memory, Request Count, 5xx, 4xx, Response Time, Healthy/Unhealthy, Custom Metrics, Running Tasks.")
todo("Insertar pantalla del dashboard con widgets")
h2("11.4 Log Metric Filters")
table(["Filter", "Pattern", "Metric"], [
    ["filter-5xx-errors", "{ $.status >= 500 }", "5xxErrors (Multicines/{env})"],
    ["filter-latency", "{ $.duration_ms = * }", "Latency (Multicines/{env})"],
])
h2("11.5 X-Ray")
para("El sidecar ADOT exporta trazas a X-Ray. Visualizar en: X-Ray → Traces → filtrar por servicio: integrator")
todo("Insertar pantalla de X-Ray Traces")
cli_box("""# Crear SNS Topic + suscripción
aws sns create-topic --name dev-multicines-integration-alarms \\
  --tags Key=Name,Value=dev-multicines-integration-alarms Key=Project,Value=multicines --region us-east-1
aws sns subscribe --topic-arn <TOPIC_ARN> --protocol email \\
  --notification-endpoint alarms@multicines.com.ec --region us-east-1

# Crear alarma (ejemplo: CPU alta)
aws cloudwatch put-metric-alarm --alarm-name dev-multicines-integration-alarm-cpu-high \\
  --namespace AWS/ECS --metric-name CPUUtilization \\
  --dimensions Name=ServiceName,Value=dev-multicines-integration-svc Name=ClusterName,Value=multicines-cluster \\
  --statistic Average --period 300 --evaluation-periods 2 --threshold 80 \\
  --comparison-operator GreaterThanOrEqualToThreshold \\
  --alarm-actions <SNS_TOPIC_ARN> --region us-east-1

# Crear dashboard
aws cloudwatch put-dashboard --dashboard-name dev-multicines-integration-dashboard \\
  --dashboard-body file://dashboard.json --region us-east-1

# Crear metric filter
aws logs put-metric-filter --log-group-name /ecs/dev-multicines-integration \\
  --filter-name dev-multicines-integration-filter-5xx-errors \\
  --filter-pattern '{ $.status >= 500 }' \\
  --metric-transformations metricName=5xxErrors,metricNamespace=Multicines/dev,metricValue=1 --region us-east-1

# IAM Policy X-Ray
aws iam put-role-policy --role-name multicines-ecs-task-role \\
  --policy-name dev-multicines-integration-xray-policy \\
  --policy-document file://xray-policy.json

# Listar alarmas
aws cloudwatch describe-alarms --alarm-name-prefix dev-multicines-integration --region us-east-1

# Eliminar alarmas
aws cloudwatch delete-alarms --alarm-names <ALARM1> <ALARM2> --region us-east-1""")

# =====================================================================
# SECTION 12: VPN
# =====================================================================
page_break()
h1("12. VPN Site-to-Site")
para("Conexión VPN IPSec entre AWS y la red on-premise de Multicines (IP: 200.7.217.58). Ambos ambientes comparten el mismo Customer Gateway.")
h2("12.1 Estado actual")
table(["Recurso", "ID", "Nombre", "Ambiente"], [
    ["VPN Connection", "vpn-08add07f10aaf549d", "dev-multicines-integration-vpn", "dev"],
    ["VPN Connection", "vpn-03ad7d548894fe350", "prod-multicines-integration-vpn", "prod"],
    ["Virtual Private GW", "vgw-0d4a9bbfacc384bdc", "dev-multicines-integration-vgw", "dev (vpc-0a5ebea8e7bee9ed5)"],
    ["Virtual Private GW", "vgw-0954e2b1a370ae78e", "prod-multicines-integration-vgw", "prod (vpc-0860a0d2cf4df6140)"],
    ["Customer Gateway", "cgw-090efcb086447cd1a", "cgw-multicines (IP: 200.7.217.58)", "compartido"],
])
h2("12.2 Customer Gateway")
steps(["VPC → Customer Gateways → Create Customer Gateway", "Name: cgw-multicines", "Routing: Static", "IP Address: 200.7.217.58 (IP pública router Multicines)", "BGP ASN: 65000 → Create"])
todo("Insertar pantalla de creación de Customer Gateway")
h2("12.3 Virtual Private Gateway")
steps(["VPC → Virtual Private Gateways → Create Virtual Private Gateway", "Name: {env}-multicines-integration-vgw → Create", "Actions → Attach to VPC → seleccionar VPC del ambiente"])
todo("Insertar pantalla del VGW adjunto a la VPC")
h2("12.4 VPN Connection")
steps(["VPC → Site-to-Site VPN Connections → Create VPN Connection", "Name: {env}-multicines-integration-vpn", "Target Gateway: VGW creado", "Customer Gateway: cgw-multicines (cgw-090efcb086447cd1a)", "Routing: Static → Create", "Download Configuration para configurar router on-premise"])
todo("Insertar pantalla del wizard VPN Connection")
h2("12.5 Route Propagation")
steps(["VPC → Route Tables → tabla privada del ambiente", "Route propagation → Edit → habilitar para VGW → Save", "Verificar ruta 192.168.0.0/16 → VGW aparece en Routes"])
todo("Insertar pantalla de route propagation")
h2("12.6 Configuración de Red")
table(["Parámetro", "Dev", "Prod"], [
    ["VPN Connection", "vpn-08add07f10aaf549d", "vpn-03ad7d548894fe350"],
    ["VGW", "vgw-0d4a9bbfacc384bdc", "vgw-0954e2b1a370ae78e"],
    ["CGW", "cgw-090efcb086447cd1a (compartido)", "cgw-090efcb086447cd1a (compartido)"],
    ["Customer IP", "200.7.217.58", "200.7.217.58"],
    ["BGP ASN", "65000", "65000"],
    ["Ruta on-premise", "192.168.0.0/16 → VGW", "192.168.0.0/16 → VGW"],
    ["VPC", "vpc-0a5ebea8e7bee9ed5", "vpc-0860a0d2cf4df6140"],
])
cli_box("""# Crear Customer Gateway (compartido)
aws ec2 create-customer-gateway --type ipsec.1 --public-ip 200.7.217.58 --bgp-asn 65000 \\
  --tag-specifications 'ResourceType=customer-gateway,Tags=[{Key=Name,Value=cgw-multicines}]' \\
  --region us-east-1

# Crear Virtual Private Gateway (por ambiente)
aws ec2 create-vpn-gateway --type ipsec.1 \\
  --tag-specifications 'ResourceType=vpn-gateway,Tags=[{Key=Name,Value=dev-multicines-integration-vgw}]' --region us-east-1
aws ec2 attach-vpn-gateway --vpn-gateway-id vgw-0d4a9bbfacc384bdc --vpc-id vpc-0a5ebea8e7bee9ed5 --region us-east-1

# Crear VPN Connection
aws ec2 create-vpn-connection --type ipsec.1 \\
  --customer-gateway-id cgw-090efcb086447cd1a --vpn-gateway-id vgw-0d4a9bbfacc384bdc \\
  --tag-specifications 'ResourceType=vpn-connection,Tags=[{Key=Name,Value=dev-multicines-integration-vpn}]' --region us-east-1

# Habilitar route propagation
aws ec2 enable-vgw-route-propagation --gateway-id vgw-0d4a9bbfacc384bdc \\
  --route-table-id <RTB_PRIVATE_ID> --region us-east-1

# Verificar estado de túneles
aws ec2 describe-vpn-connections --vpn-connection-ids vpn-08add07f10aaf549d \\
  --query "VpnConnections[0].VgwTelemetry[].{Status:Status,IP:OutsideIpAddress}" --region us-east-1

# Actualizar tags
aws ec2 create-tags --resources vpn-08add07f10aaf549d \\
  --tags Key=Name,Value=dev-multicines-integration-vpn Key=Project,Value=multicines Key=Environment,Value=dev --region us-east-1""")

# =====================================================================
# SECTION 13: CI/CD
# =====================================================================
page_break()
h1("13. CI/CD — GitHub Actions")
h2("13.1 Proveedor OIDC")
steps(["IAM → Identity providers → Add provider", "Type: OpenID Connect", "URL: https://token.actions.githubusercontent.com", "Audience: sts.amazonaws.com → Add provider"])
add_image("image4c.png")
add_image("image4d.png")
add_image("image4e.png")
h2("13.2 Rol IAM GitHub Actions")
para("Rol: multicines-github-actions-role | ARN: arn:aws:iam::340271092920:role/multicines-github-actions-role")
para("Políticas: AmazonECS_FullAccess + AmazonEC2ContainerRegistryPowerUser")
add_image("image4f.png")
add_image("image50.png")
add_image("image51.png")
add_image("image52.png")
add_image("image53.png")
add_image("image54.png")
h2("13.3 Variables y Secrets en GitHub")
table(["Variable", "Valor"], [
    ["AWS_REGION", "us-east-1"],
    ["ECR_REPOSITORY", "multicines/integration-service"],
    ["ECS_CLUSTER", "multicines-cluster"],
    ["DEV_TASK_DEFINITION", "dev-multicines-integration-task"],
    ["PROD_TASK_DEFINITION", "prod-multicines-integration-task"],
    ["DEV_ECS_SERVICE", "dev-multicines-integration-svc"],
    ["PROD_ECS_SERVICE", "prod-multicines-integration-svc"],
])
table(["Secret", "Valor"], [
    ["AWS_ROLE_ARN", "arn:aws:iam::340271092920:role/multicines-github-actions-role"],
])
h2("13.4 Pipeline")
table(["Workflow", "Trigger", "Ambiente"], [
    ["deploy-dev.yml", "Push a dev", "development"],
    ["deploy-prod.yml", "Push a main", "production"],
])
para("Flujo: Checkout → OIDC Auth → ECR Login → Build Docker → Push → Update TD → Deploy ECS → Verify → (On failure: Rollback)")
h2("13.5 Protección de ramas")
add_image("image55.png")
add_image("image56.png")
add_image("image57.png")
add_image("image58.png")
add_image("image59.png")
add_image("image5a.png")
add_image("image5b.png")
cli_box("""# Crear OIDC provider
aws iam create-open-id-connect-provider \\
  --url https://token.actions.githubusercontent.com \\
  --client-id-list sts.amazonaws.com

# Crear rol GitHub Actions
aws iam create-role --role-name multicines-github-actions-role \\
  --assume-role-policy-document file://github-trust-policy.json
aws iam attach-role-policy --role-name multicines-github-actions-role \\
  --policy-arn arn:aws:iam::aws:policy/AmazonECS_FullAccess
aws iam attach-role-policy --role-name multicines-github-actions-role \\
  --policy-arn arn:aws:iam::aws:policy/AmazonEC2ContainerRegistryPowerUser""")

# === SAVE ===
doc.save(OUTPUT)
sz = os.path.getsize(OUTPUT)
print(f"\n[DONE] {OUTPUT}")
print(f"[DONE] Size: {sz / 1024:.0f} KB")
