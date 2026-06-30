"""
Actualiza el Manual Multicines v2.docx agregando secciones nuevas al final.
Preserva todo el contenido y formato existente del documento original.
"""
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

DOCX_PATH = os.path.join(os.path.dirname(__file__), "Manual Multicines v2.docx")

doc = Document(DOCX_PATH)

# === Helper functions ===
def add_heading1(text):
    p = doc.add_heading(text, level=1)
    return p

def add_heading2(text):
    p = doc.add_heading(text, level=2)
    return p

def add_heading3(text):
    p = doc.add_heading(text, level=3)
    return p

def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p

def add_note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return p

def add_todo(text):
    p = doc.add_paragraph()
    run = p.add_run(f"[TODO: {text}]")
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = str(val)
    doc.add_paragraph()  # spacing
    return table

def add_numbered_steps(steps):
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph(f"{i}. {step}")
    doc.add_paragraph()

# === PAGE BREAK before new content ===
doc.add_page_break()

# =============================================================================
# SECTION: ACM Certificate
# =============================================================================
add_heading1("5. Certificados SSL — ACM")

add_heading2("5.1 Importación del certificado")
add_para("Se importa un certificado SSL/TLS wildcard (*.test.multicines.com.ec) en AWS Certificate Manager para la terminación HTTPS en el ALB. Al ser un certificado importado, requiere renovación manual antes de su expiración.")

add_heading3("Paso a paso en consola")
add_numbered_steps([
    "Buscar Certificate Manager en la barra de búsqueda de la consola AWS",
    "Click en Import certificate",
    "En Certificate body: pegar contenido de certificate.pem",
    "En Certificate private key: pegar contenido de private-key.pem",
    "En Certificate chain: pegar contenido de certificate-chain.pem",
    "Click Next → agregar tags: Name={env}-multicines-integration-acm, Project=multicines, Environment={env}",
    "Click Import"
])
add_todo("Insertar pantalla del formulario de importación de certificado ACM")
add_todo("Insertar pantalla del certificado importado mostrando estado Issued")

add_heading3("Parámetros de configuración")
add_table(
    ["Parámetro", "Dev", "Prod"],
    [
        ["Dominio", "*.test.multicines.com.ec", "*.test.multicines.com.ec"],
        ["Tipo", "Imported", "Imported"],
        ["Tag Name", "dev-multicines-integration-acm", "prod-multicines-integration-acm"],
        ["Archivos fuente", "environments/dev/certs/", "environments/prod/certs/"],
        ["Recurso asociado", "ALB listener HTTPS:443", "ALB listener HTTPS:443"],
    ]
)

add_heading3("Renovación del certificado")
add_para("Al ser un certificado importado, AWS no lo renueva automáticamente. Para renovar:")
add_numbered_steps([
    "Obtener nuevo certificado del proveedor SSL",
    "En ACM → seleccionar certificado existente → Reimport certificate",
    "Pegar los nuevos archivos PEM",
    "Click Import — el ARN no cambia, no requiere re-asociar al ALB"
])

# =============================================================================
# SECTION: WAF
# =============================================================================
doc.add_page_break()
add_heading1("6. WAF — Web Application Firewall")

add_heading2("6.1 Creación de la Web ACL")
add_para("Se crea una Web ACL regional asociada al ALB con una regla rate-based que bloquea IPs que excedan 2000 requests en 5 minutos.")

add_heading3("Paso a paso en consola")
add_numbered_steps([
    "Buscar WAF & Shield → Web ACLs → Create web ACL",
    "Resource type: Regional resources",
    "Name: {env}-multicines-integration-waf",
    "Region: US East (N. Virginia)",
    "Associated AWS resources → Add AWS resources → seleccionar el ALB del ambiente",
    "Click Next",
    "Add rules → Add my own rules → Rule builder",
    "Rule type: Rate-based rule",
    "Name: rate-limit-per-ip",
    "Rate limit: 2000 requests por 5 minutos",
    "IP address to use: Source IP address",
    "Action: Block",
    "Click Add rule → Next",
    "Default action: Allow",
    "Completar wizard → Create web ACL"
])
add_todo("Insertar pantalla del wizard WAF (paso nombre y asociación con ALB)")
add_todo("Insertar pantalla de la regla rate-based configurada")

add_heading3("Parámetros de configuración")
add_table(
    ["Parámetro", "Dev", "Prod"],
    [
        ["Nombre Web ACL", "dev-multicines-integration-waf", "prod-multicines-integration-waf"],
        ["Scope", "REGIONAL", "REGIONAL"],
        ["Default action", "Allow", "Allow"],
        ["Rate limit", "2000 req/5min por IP", "2000 req/5min por IP"],
        ["Rate limit action", "Block", "Block"],
        ["Recurso asociado", "dev-multicines-integration-alb", "prod-multicines-integration-alb"],
    ]
)

# =============================================================================
# SECTION: Security Groups Hardening
# =============================================================================
doc.add_page_break()
add_heading1("7. Security Groups — Hardening de Egress")

add_para("Los Security Groups ya existen. Este paso modifica las reglas de salida (egress) para restringir el tráfico saliente al mínimo necesario.")

add_heading2("7.1 SG del ALB")
add_numbered_steps([
    "VPC → Security Groups → seleccionar SG ALB (sg-0fdf0c744482646ae dev / sg-04f21c4c1d065fe60 prod)",
    "Pestaña Outbound rules → Edit outbound rules",
    "Eliminar la regla All traffic → 0.0.0.0/0",
    "Add rule: Type=Custom TCP, Port=8095, Destination=SG ECS del ambiente",
    "Click Save rules"
])
add_todo("Insertar pantalla de Outbound rules del SG ALB después del hardening")

add_heading2("7.2 SG de ECS")
add_numbered_steps([
    "VPC → Security Groups → seleccionar SG ECS (sg-01c619444bab0b8d3 dev / sg-0b78c14cd02b2f6e1 prod)",
    "Pestaña Outbound rules → Edit outbound rules",
    "Eliminar la regla All traffic → 0.0.0.0/0",
    "Add rule: Type=HTTPS, Port=443, Destination=0.0.0.0/0",
    "Click Save rules"
])
add_todo("Insertar pantalla de Outbound rules del SG ECS después del hardening")

add_heading3("Resultado esperado")
add_table(
    ["SG", "Egress permitido", "Propósito"],
    [
        ["ALB", "TCP 8095 → SG ECS", "Solo tráfico al contenedor"],
        ["ECS", "TCP 443 → 0.0.0.0/0", "Servicios AWS (ECR, CloudWatch, SSM, Secrets)"],
    ]
)

# =============================================================================
# SECTION: ALB
# =============================================================================
doc.add_page_break()
add_heading1("8. Application Load Balancer")

add_heading2("8.1 Creación del ALB")
add_numbered_steps([
    "EC2 → Load Balancers → Create Load Balancer → Application Load Balancer",
    "Name: {env}-multicines-integration-alb",
    "Scheme: Internet-facing, IP type: IPv4",
    "Network mapping: seleccionar VPC y las 2 subnets públicas del ambiente",
    "Security groups: seleccionar SG ALB",
    "Listeners: HTTPS:443 → Forward to Target Group",
    "Seleccionar certificado ACM *.test.multicines.com.ec",
    "SSL Policy: ELBSecurityPolicy-TLS13-1-2-2021-06",
    "Click Create load balancer"
])
add_todo("Insertar pantalla del wizard de creación del ALB")

add_heading2("8.2 Creación del Target Group")
add_numbered_steps([
    "EC2 → Target Groups → Create target group",
    "Target type: IP addresses (para Fargate)",
    "Name: {env}-multicines-integration-tg",
    "Protocol: HTTP, Port: 8095",
    "VPC: seleccionar la del ambiente",
    "Health check path: /actuator/health, Port: 8095",
    "Healthy threshold: 2, Unhealthy threshold: 5, Timeout: 10s, Interval: 30s",
    "Click Create target group"
])

add_heading3("Parámetros del ALB")
add_table(
    ["Parámetro", "Dev", "Prod"],
    [
        ["Nombre ALB", "dev-multicines-integration-alb", "prod-multicines-integration-alb"],
        ["Tipo", "application, internet-facing", "application, internet-facing"],
        ["SSL Policy", "ELBSecurityPolicy-TLS13-1-2-2021-06", "ELBSecurityPolicy-TLS13-1-2-2021-06"],
        ["Target Group", "dev-multicines-integration-tg", "prod-multicines-integration-tg"],
        ["Health Check Path", "/actuator/health", "/actuator/health"],
        ["Health Check Port", "8095", "8095"],
    ]
)

add_heading2("8.3 Listener HTTP → HTTPS redirect")
add_numbered_steps([
    "En el ALB → Pestaña Listeners → Add listener",
    "Protocol: HTTP, Port: 80",
    "Default action: Redirect to URL → HTTPS:443, Status code: 301",
    "Click Add"
])

# =============================================================================
# SECTION: Route53
# =============================================================================
doc.add_page_break()
add_heading1("9. Route 53 — DNS")

add_heading2("9.1 Creación de la Hosted Zone")
add_numbered_steps([
    "Buscar Route 53 → Hosted zones → Create hosted zone",
    "Domain name: test.multicines.com.ec",
    "Type: Public hosted zone",
    "Click Create hosted zone",
    "IMPORTANTE: Copiar los 4 nameservers (NS records) y configurarlos en el registrador del dominio"
])
add_todo("Insertar pantalla de la hosted zone con NS records")

add_heading2("9.2 Registro DNS Alias al ALB")
add_numbered_steps([
    "Dentro de la hosted zone → Create record",
    "Record name: dev-integration (o prod-integration)",
    "Record type: A",
    "Toggle Alias: ON",
    "Route traffic to: Alias to Application and Classic Load Balancer",
    "Region: US East (N. Virginia)",
    "Seleccionar el ALB correspondiente",
    "Evaluate target health: Yes",
    "Click Create records"
])
add_todo("Insertar pantalla del formulario de creación de record alias A")

add_heading3("Registros DNS configurados")
add_table(
    ["Registro", "Target", "Ambiente"],
    [
        ["dev-integration.test.multicines.com.ec", "ALB dev", "dev"],
        ["prod-integration.test.multicines.com.ec", "ALB prod", "prod"],
    ]
)

# =============================================================================
# SECTION: SSM
# =============================================================================
doc.add_page_break()
add_heading1("10. SSM Parameter Store")

add_heading2("10.1 Creación de parámetros")
add_para("Los parámetros almacenan configuración sensible cifrada con KMS bajo la jerarquía /multicines/integrator/{env}/.")

add_numbered_steps([
    "Buscar Systems Manager → panel izquierdo → Parameter Store",
    "Click Create parameter",
    "Name: /multicines/integrator/{env}/{nombre_parametro}",
    "Tier: Standard",
    "Type: SecureString",
    "KMS Key Source: My current account → aws/ssm (default)",
    "Value: pegar el valor JSON compactado",
    "Tags: Name, Project=multicines, Environment={env}, Service=multicines-integration",
    "Click Create parameter",
    "Repetir para cada parámetro requerido"
])
add_todo("Insertar pantalla del formulario de creación de parámetro SSM")

add_heading3("Parámetro fijo OTEL")
add_para("Adicionalmente crear un parámetro de tipo String (no cifrado):")
add_table(
    ["Name", "Type", "Value"],
    [
        ["/multicines/integrator/{env}/OTEL_TRACES_EXPORTER", "String", "otlp"],
    ]
)

# =============================================================================
# SECTION: Secrets Manager
# =============================================================================
doc.add_page_break()
add_heading1("11. Secrets Manager")

add_heading2("11.1 Creación de secretos")
add_numbered_steps([
    "Buscar Secrets Manager → Store a new secret",
    "Secret type: Other type of secret",
    "Key/value: seleccionar Plaintext → pegar el JSON compactado",
    "Encryption key: aws/secretsmanager (default)",
    "Click Next",
    "Secret name: escribir la key completa (ej: platform/dev/resilience)",
    "Tags: Name={key}, Project=multicines, Environment={env}, Service=multicines-integration",
    "Click Next → Skip rotation → Store",
    "Repetir para cada secreto requerido"
])
add_todo("Insertar pantalla del wizard de creación de secreto")

# =============================================================================
# SECTION: Observabilidad
# =============================================================================
doc.add_page_break()
add_heading1("15. Observabilidad — CloudWatch, X-Ray")

add_heading2("15.1 SNS Topic para notificaciones")
add_numbered_steps([
    "Buscar SNS → Topics → Create topic",
    "Type: Standard",
    "Name: {env}-multicines-integration-alarms",
    "Tags: Name, Project=multicines, Environment={env}",
    "Click Create topic",
    "En el topic → Create subscription → Protocol: Email → Endpoint: alarms@multicines.com.ec",
    "Click Create subscription",
    "Confirmar la suscripción desde el email recibido"
])
add_todo("Insertar pantalla del SNS Topic con suscripción email")

add_heading2("15.2 CloudWatch Alarms")
add_para("Se crean 8 alarmas que notifican al SNS Topic cuando se superan umbrales críticos.")

add_heading3("Creación de una alarma (ejemplo: CPU alta)")
add_numbered_steps([
    "CloudWatch → Alarms → Create alarm",
    "Select metric → AWS/ECS → ClusterName, ServiceName → CPUUtilization",
    "Seleccionar servicio {env}-multicines-integration-svc en multicines-cluster",
    "Statistic: Average, Period: 5 minutes",
    "Condition: Greater/Equal than → 80",
    "Click Next → Notification: In alarm → SNS topic {env}-multicines-integration-alarms",
    "Alarm name: {env}-multicines-integration-alarm-cpu-high",
    "Click Create alarm"
])
add_todo("Insertar pantalla del wizard de creación de alarma")

add_heading3("Alarmas configuradas")
add_table(
    ["Nombre", "Métrica", "Threshold", "Período"],
    [
        ["alarm-cpu-high", "AWS/ECS CPUUtilization", "≥ 80%", "300s, 2 eval"],
        ["alarm-memory-high", "AWS/ECS MemoryUtilization", "≥ 80%", "300s, 2 eval"],
        ["alarm-5xx", "ALB HTTPCode_Target_5XX_Count", "≥ 10", "60s, 3 eval"],
        ["alarm-4xx", "ALB HTTPCode_Target_4XX_Count", "≥ 50", "60s, 3 eval"],
        ["alarm-unhealthy-hosts", "ALB UnHealthyHostCount", "≥ 1", "60s, 2 eval"],
        ["alarm-response-time-high", "ALB TargetResponseTime", "≥ 5s/3s", "60s, 3 eval"],
        ["alarm-no-running-tasks", "ECS RunningTaskCount", "< 1", "60s, 2 eval"],
        ["alarm-vpn-tunnel-down", "VPN TunnelState", "< 1", "60s, 2 eval"],
    ]
)
add_note("Nota: La alarma VPN solo se crea cuando la IP del Customer Gateway no es placeholder (0.0.0.0).")

add_heading2("15.3 Log Metric Filters")
add_numbered_steps([
    "CloudWatch → Logs → Log groups → /ecs/{env}-multicines-integration",
    "Pestaña Metric filters → Create metric filter",
    "Filter pattern: { $.status >= 500 }",
    "Click Test pattern → Click Next",
    "Filter name: {env}-multicines-integration-filter-5xx-errors",
    "Metric namespace: Multicines/{env}, Metric name: 5xxErrors, Metric value: 1",
    "Click Create metric filter"
])
add_para("Repetir para el filtro de latencia con pattern: { $.duration_ms = * } y metric value: $.duration_ms")
add_todo("Insertar pantalla del formulario de metric filter")

add_heading2("15.4 CloudWatch Dashboard")
add_numbered_steps([
    "CloudWatch → Dashboards → Create dashboard",
    "Name: {env}-multicines-integration-dashboard",
    "Agregar widgets: CPU, Memory, Request Count, 5xx, 4xx, Response Time, Healthy/Unhealthy Hosts, Custom Metrics, Running Tasks, VPN Tunnel",
    "Para cada widget: Add widget → Line → seleccionar métrica → período 5 min → Create widget",
    "Organizar layout → Save dashboard"
])
add_todo("Insertar pantalla del dashboard completo con todos los widgets")

add_heading2("15.5 IAM Policy X-Ray")
add_numbered_steps([
    "IAM → Roles → click en multicines-ecs-task-role",
    "Pestaña Permissions → Add permissions → Create inline policy",
    "Click pestaña JSON → pegar policy con acciones: xray:PutTraceSegments, xray:PutTelemetryRecords, xray:GetSamplingRules, xray:GetSamplingTargets",
    "Resource: *",
    "Click Next → Policy name: {env}-multicines-integration-xray-policy",
    "Click Create policy"
])
add_todo("Insertar pantalla de la inline policy X-Ray en el rol")

add_heading2("15.6 Verificación de X-Ray")
add_numbered_steps([
    "Buscar X-Ray → Traces",
    "Filtrar por servicio: integrator",
    "Click en una traza para ver el flujo distribuido (latencia por segmento)"
])
add_todo("Insertar pantalla de X-Ray Traces del servicio integrator")

# === SAVE ===
doc.save(DOCX_PATH)
print(f"[OK] Documento actualizado: {DOCX_PATH}")
print(f"[OK] Tamaño: {os.path.getsize(DOCX_PATH) / 1024:.0f} KB")
