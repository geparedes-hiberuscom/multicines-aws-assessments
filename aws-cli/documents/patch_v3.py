"""
Patch v3: Update 1.4 tables in-place + remove Anexo.
Approach: find existing tables in 1.4 section and replace their content.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

DOCX_PATH = os.path.join(os.path.dirname(__file__), "Manual de Gesti�n de Recursos en AWS v1.0.docx")
doc = Document(DOCX_PATH)
body = doc.element.body

# === Find the 3 tables in section 1.4 ===
# They are between paragraph 18 (1.4 heading) and 26 (section 2)
# Find table elements in that range
h14_elem = doc.paragraphs[18]._element
h2_elem = doc.paragraphs[26]._element

tables_in_14 = []
for tbl in doc.tables:
    tbl_elem = tbl._tbl
    # Check if table is between 1.4 and section 2
    prev = tbl_elem.getprevious()
    parent = tbl_elem.getparent()
    if parent is body:
        # Check position relative to our markers
        all_body = list(body)
        tbl_idx = all_body.index(tbl_elem)
        h14_body_idx = all_body.index(h14_elem)
        h2_body_idx = all_body.index(h2_elem)
        if h14_body_idx < tbl_idx < h2_body_idx:
            tables_in_14.append(tbl)

print(f"[INFO] Found {len(tables_in_14)} tables in section 1.4")


# === Data for the 3 tables ===
shared_data = [
    ["Route53", "Hosted Zone", "cloudmulticines.com (Z02656981PPWUEVYOEMVL)"],
    ["ACM", "Certificado SSL", "*.cloudmulticines.com (ISSUED)"],
    ["ECR", "Repositorio", "multicines/integration-service"],
    ["ECS", "Cluster", "multicines-cluster"],
    ["IAM", "Execution Role", "ecsTaskExecutionRole"],
    ["IAM", "Task Role", "multicines-ecs-task-role"],
    ["IAM", "GitHub Actions Role", "multicines-github-actions-role"],
    ["EC2", "Customer Gateway", "cgw-090efcb086447cd1a (IP: 200.7.217.58)"],
]

dev_data = [
    ["VPC", "VPC", "vpc-0a5ebea8e7bee9ed5"],
    ["VPC", "Subnet pública A", "subnet-0c92479c2831f04eb"],
    ["VPC", "Subnet pública B", "subnet-02193c97ceb6ec25d"],
    ["VPC", "Subnet privada", "subnet-0eeec60453f7f9492"],
    ["EC2", "SG ALB", "sg-0fdf0c744482646ae"],
    ["EC2", "SG ECS", "sg-01c619444bab0b8d3"],
    ["ALB", "Load Balancer", "dev-multicines-integration-alb"],
    ["ALB", "Target Group", "dev-multicines-integration-tg (port 8095)"],
    ["Route53", "Record A", "api-dev.cloudmulticines.com"],
    ["ECS", "Service", "dev-multicines-integration-svc (1 task)"],
    ["ECS", "Task Definition", "dev-multicines-integration-task"],
    ["Secrets", "Secret", "platform/dev/resilience"],
    ["Secrets", "Secret", "integrator/dev/app"],
    ["SNS", "Topic", "dev-multicines-integration-alarms"],
    ["CloudWatch", "Dashboard", "dev-multicines-integration-dashboard"],
    ["CloudWatch", "Alarms", "7 alarmas activas"],
    ["VPN", "Connection", "vpn-08add07f10aaf549d"],
    ["VPN", "Virtual Private GW", "vgw-0d4a9bbfacc384bdc"],
]

prod_data = [
    ["VPC", "VPC", "vpc-0860a0d2cf4df6140"],
    ["VPC", "Subnet pública A", "subnet-04e4cd18763fd6f84"],
    ["VPC", "Subnet pública B", "subnet-05c92a4889326c68e"],
    ["VPC", "Subnet privada A", "subnet-0d55ebfe5adbd4e81"],
    ["VPC", "Subnet privada B", "subnet-04f8ac3335c426f57"],
    ["EC2", "SG ALB", "sg-04f21c4c1d065fe60"],
    ["EC2", "SG ECS", "sg-0b78c14cd02b2f6e1"],
    ["ALB", "Load Balancer", "prod-multicines-integration-alb"],
    ["ALB", "Target Group", "prod-multicines-integration-tg (port 8095)"],
    ["Route53", "Record A", "api.cloudmulticines.com"],
    ["ECS", "Service", "prod-multicines-integration-svc (2 tasks)"],
    ["ECS", "Task Definition", "prod-multicines-integration-task"],
    ["Secrets", "Secret", "platform/prod/resilience"],
    ["Secrets", "Secret", "integrator/prod/app"],
    ["SNS", "Topic", "prod-multicines-integration-alarms"],
    ["CloudWatch", "Dashboard", "prod-multicines-integration-dashboard"],
    ["VPN", "Connection", "vpn-03ad7d548894fe350"],
    ["VPN", "Virtual Private GW", "vgw-0954e2b1a370ae78e"],
]

datasets = [shared_data, dev_data, prod_data]
headers = ["Servicio", "Recurso", "ID/Nombre"]

# === Update tables in place ===
def rebuild_table(tbl, hdr, data):
    """Clear table and rebuild with new data"""
    # Remove all rows except header
    while len(tbl.rows) > 1:
        tr = tbl.rows[-1]._tr
        tbl._tbl.remove(tr)
    
    # Update header
    for i, h in enumerate(hdr):
        cell = tbl.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        # Remove existing shading if any
        existing = tcPr.find(qn('w:shd'))
        if existing is not None:
            tcPr.remove(existing)
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F4E79" w:val="clear"/>'))
    
    # Add data rows
    for row_data in data:
        row = tbl.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.text = ''
            r = cell.paragraphs[0].add_run(str(val))
            r.font.size = Pt(9)

for i, tbl in enumerate(tables_in_14):
    if i < len(datasets):
        rebuild_table(tbl, headers, datasets[i])
        print(f"  Updated table {i+1} with {len(datasets[i])} rows")

# === STEP 2: Remove Anexo (section 14) ===
print("[INFO] Removing Anexo section...")
anexo_elem = None
for p in doc.paragraphs:
    if p.style.name == 'Heading 1' and 'Anexo' in p.text:
        anexo_elem = p._element
        break

if anexo_elem is not None:
    all_elems = list(body)
    removing = False
    removed = 0
    for elem in all_elems:
        if elem is anexo_elem:
            removing = True
        if removing and not elem.tag.endswith('}sectPr'):
            body.remove(elem)
            removed += 1
    print(f"  Removed {removed} elements")
else:
    print("  Anexo not found")

# === SAVE ===
doc.save(DOCX_PATH)
print(f"\n[DONE] {DOCX_PATH}")
print(f"[DONE] Size: {os.path.getsize(DOCX_PATH) / 1024:.0f} KB")
